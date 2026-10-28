"""
Generates a polished 8-Slide PowerPoint Presentation for the Greenhouse Monitoring Project.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
import os

# Colors
NAVY = RGBColor(15, 23, 42)       # #0F172A
SLATE = RGBColor(30, 41, 59)      # #1E293B
GREEN = RGBColor(16, 185, 129)    # #10B981
EMERALD = RGBColor(5, 150, 105)   # #059669
WHITE = RGBColor(255, 255, 255)
LIGHT_BG = RGBColor(248, 250, 252)# #F8FAFC
DARK_GRAY = RGBColor(51, 65, 85)  # #334155
MUTED = RGBColor(100, 116, 139)   # #64748B
ACCENT_BLUE = RGBColor(37, 99, 235)


def create_deck():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    def add_header(slide, title_text, step_badge=""):
        # Header background bar
        header_box = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.15)
        )
        header_box.fill.solid()
        header_box.fill.fore_color.rgb = NAVY
        header_box.line.fill.background()

        # Title text
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.8)
        tf.margin_top = Inches(0.2)
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.font.name = "Arial"

        # Badge if present
        if step_badge:
            badge_box = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.2), Inches(0.3), Inches(2.3), Inches(0.55)
            )
            badge_box.fill.solid()
            badge_box.fill.fore_color.rgb = GREEN
            badge_box.line.fill.background()
            btf = badge_box.text_frame
            bp = btf.paragraphs[0]
            bp.text = step_badge
            bp.font.size = Pt(13)
            bp.font.bold = True
            bp.font.color.rgb = WHITE
            bp.font.name = "Arial"
            bp.alignment = PP_ALIGN.CENTER

    def add_card(slide, left, top, width, height, title, items, color=SLATE, bg_color=WHITE):
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = RGBColor(226, 232, 240)
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.35)
        tf.margin_right = Inches(0.35)
        tf.margin_top = Inches(0.3)

        # Title
        p_title = tf.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(18)
        p_title.font.bold = True
        p_title.font.color.rgb = color
        p_title.font.name = "Arial"

        # Items
        for item in items:
            p = tf.add_paragraph()
            p.text = "•  " + item
            p.font.size = Pt(14)
            p.font.color.rgb = DARK_GRAY
            p.font.name = "Arial"
            p.space_before = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide (Dark Theme)
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = NAVY
    bg1.line.fill.background()

    # Title Box
    t_box = slide1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.333), Inches(3.5))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p1 = tf1.paragraphs[0]
    p1.text = "🌱 SMART GREENHOUSE MONITORING SYSTEM"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = GREEN
    p1.font.name = "Arial"

    p2 = tf1.add_paragraph()
    p2.text = "Automated Real-Time Telemetry, Plant Analytics & Closed-Loop Anomaly Control"
    p2.font.size = Pt(20)
    p2.font.color.rgb = WHITE
    p2.font.name = "Arial"
    p2.space_before = Pt(14)

    p3 = tf1.add_paragraph()
    p3.text = "Flow: Sensors ➔ Database ➔ Data Analytics ➔ Anomaly Detection ➔ Hardware Output"
    p3.font.size = Pt(15)
    p3.font.color.rgb = RGBColor(148, 163, 184)
    p3.font.name = "Arial"
    p3.space_before = Pt(28)

    # -------------------------------------------------------------
    # SLIDE 2: Problem Statement & 5-Stage Architecture
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "Project Overview & 5-Stage Architecture", "Architecture")

    add_card(slide2, 0.8, 1.6, 5.5, 5.2, "⚠️ The Horticultural Challenge", [
        "Uncontrolled climate swings (heatwaves, cold snaps) severely stunt plant growth.",
        "Excessive humidity fosters mold, mildew, and fungal outbreaks.",
        "Stale air and CO2 depletion stall photosynthesis, reducing crop yield.",
        "Manual inspection is slow, labor-intensive, and prone to human error.",
        "Need: An automated 24/7 closed-loop monitoring & protection system."
    ], color=RGBColor(225, 29, 72))

    add_card(slide2, 6.8, 1.6, 5.7, 5.2, "⚡ The 5-Stage Solution Flow", [
        "Stage 1 (Sensors): Physical acquisition via ESP32, DHT22, LDR, MQ-135.",
        "Stage 2 (Database): Ingestion via Flask REST API into SQLite database.",
        "Stage 3 (Analytics): Calculating VPD, CO2 availability, and Oxygen balance.",
        "Stage 4 (Anomalies): Multi-parameter rule checks and severity scoring.",
        "Stage 5 (Hardware Output): Closed-loop relay triggers and live Streamlit UI."
    ], color=EMERALD)

    # -------------------------------------------------------------
    # SLIDE 3: Stage 1 — Sensor Data Acquisition
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "Stage 1: Sensor Data Acquisition (Edge Layer)", "Stage 1: Sensors")

    add_card(slide3, 0.8, 1.6, 5.5, 5.2, "🔌 Hardware Sensors & Pinouts", [
        "DHT22 / DHT11 (GPIO 4): Temperature (°C) & Relative Humidity (%).",
        "LDR Light Sensor (GPIO 34): Solar light converted to Lux (lx).",
        "MQ-135 Gas Sensor (GPIO 35): Air quality converted to AQI (Index).",
        "12-bit ADC (0 - 4095 counts): High-resolution analog voltage digitization.",
        "Onboard Status LED (GPIO 2): Real-time WiFi & transmission indicator."
    ], color=ACCENT_BLUE)

    add_card(slide3, 6.8, 1.6, 5.7, 5.2, "📡 Signal Processing & WiFi Transmission", [
        "Hardware Signal Smoothing: Multi-sample averaging reduces noise.",
        "Unit Conversion Engine: Non-linear mapping from ADC to Lux and AQI.",
        "Structured JSON Payload: Formatted with temperature, humidity, light, gas.",
        "HTTP POST /data: Transmitted to Flask backend over WiFi every 2-3s.",
        "Auto-reconnection logic: Ensures continuous telemetry streaming."
    ], color=GREEN)

    # -------------------------------------------------------------
    # SLIDE 4: Stage 2 — Ingestion & Database Management
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "Stage 2: Ingestion & Database Management", "Stage 2: Database")

    add_card(slide4, 0.8, 1.6, 5.5, 5.2, "🚀 Flask REST API Backend", [
        "Listens on port 5000: Handles high-throughput sensor telemetry.",
        "Endpoint POST /data: Ingests packets and returns immediate anomaly state.",
        "Query Endpoints: /latest, /history, /stats, /anomalies, and /health.",
        "CORS Enabled: Supports direct integration with dashboards & web apps.",
        "Thread-Safe Connection Pooling: Robust multi-client query handling."
    ], color=SLATE)

    add_card(slide4, 6.8, 1.6, 5.7, 5.2, "💾 ACID-Compliant SQLite Database", [
        "Database File: backend/greenhouse.db with structured table sensor_data.",
        "Schema: (id, timestamp, temperature, humidity, light, gas).",
        "High-Precision Timestamps: ISO-8601 UTC timestamps on every row.",
        "Fast B-Tree Indexing: Instant time-series lookups and chronological ordering.",
        "Historical Integrity: Clean persistent storage for offline data analytics."
    ], color=EMERALD)

    # -------------------------------------------------------------
    # SLIDE 5: Stage 3 — Data Analytics & Plant Physiology
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "Stage 3: Data Analytics & Plant Physiology", "Stage 3: Analytics")

    add_card(slide5, 0.8, 1.6, 5.5, 5.2, "🌿 Botanical Feature Engineering", [
        "Vapor Pressure Deficit (VPD in kPa): Quantifies plant transpiration.",
        "Optimal Sweet-Spot (0.8 - 1.4 kPa): Maximum photosynthesis without wilting.",
        "CO2 Availability (ppm): Tracks carbon assimilation for crop biomass.",
        "Oxygen (O2 %) Dynamic Balance: Balances day photosynthesis & gas displacement.",
        "Dew Point Temperature (°C): Predicts leaf condensation and mold hazard."
    ], color=EMERALD)

    add_card(slide5, 6.8, 1.6, 5.7, 5.2, "📈 Statistical Insights & Trends", [
        "ETL Pipeline (prepare_data.py): Automatic cleaning and physical sanity bounds.",
        "Pearson Correlation Analysis: Discovers temp-humidity inverse relationships.",
        "Descriptive Statistics: Mean, std deviation, min/max, quartiles across time.",
        "Predictive Forecasting (Holt's Method): Projects trend 15-30 min ahead.",
        "Early Warning: Warns operators before environmental setpoints breach."
    ], color=ACCENT_BLUE)

    # -------------------------------------------------------------
    # SLIDE 6: Stage 4 — Anomaly Detection Engine
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "Stage 4: Intelligent Anomaly Detection", "Stage 4: Anomalies")

    add_card(slide6, 0.8, 1.6, 5.5, 5.2, "🚨 Multi-Parameter Rule Checks", [
        "High Temperature Alert: Flags overheating conditions (Temp > 40°C).",
        "Air Quality / Smoke Spike: Flags hazardous gas accumulation (AQI > 150).",
        "Compound Agricultural Rules:",
        "  - Heat Stress + Wilt: High Temp (> 35°C) + Dry Air (< 30%).",
        "  - Fungal Outbreak: Warm Air (> 24°C) + High Saturation (> 85%).",
        "  - Leaf Scorch: Excessive Solar Irradiance (> 55,000 lx) + High Temp."
    ], color=RGBColor(225, 29, 72))

    add_card(slide6, 6.8, 1.6, 5.7, 5.2, "🧠 Severity & Actionable Remediation", [
        "Graded Severity: NORMAL (0/100) vs ANOMALY (60-100/100).",
        "Multivariate Mahalanobis ML Scoring: Detects subtle multi-sensor drift.",
        "Contextual Operator Checklists:",
        "  - 'Activate evaporative cooling and misting nozzles.'",
        "  - 'Turn on emergency exhaust fans to clear hazardous fumes.'",
        "  - 'Deploy shade screens to protect foliage from photo-inhibition.'"
    ], color=DARK_GRAY)

    # -------------------------------------------------------------
    # SLIDE 7: Stage 5 — Hardware Output, Actuation & Dashboard
    # -------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "Stage 5: Hardware Output & Real-Time Dashboard", "Stage 5: Output")

    add_card(slide7, 0.8, 1.6, 5.5, 5.2, "⚡ Closed-Loop Edge Actuation", [
        "HTTP Response Feedback: Backend sends anomaly flags back to ESP32.",
        "Relay Module Triggers: Automatically switches ON cooling fans or misting pump.",
        "Exhaust Ventilation: Purges stale air upon detecting CO2/AQI spikes.",
        "Visual & Audible Alarms: Hardware status LED and buzzers alert on-site staff.",
        "Automated 24/7 Protection: No human supervision required for rapid response."
    ], color=EMERALD)

    add_card(slide7, 6.8, 1.6, 5.7, 5.2, "🖥️ Streamlit Monitoring Dashboard", [
        "Live Telemetry Grid: Temp (°C), Humidity (%), Lux (lx), AQI, CO2, O2, VPD.",
        "Threshold Setters (Sidebar): Customizable slider limits for all metrics.",
        "Interactive Altair Trend Charts: Area gradients & dashed threshold lines.",
        "Live Condition Cards: Color-coded health banners and action checklists.",
        "CSV Data Export: One-click historical data and anomaly report download."
    ], color=ACCENT_BLUE)

    # -------------------------------------------------------------
    # SLIDE 8: Results & Conclusion
    # -------------------------------------------------------------
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(slide8, "Project Summary & Impact", "Conclusion")

    add_card(slide8, 0.8, 1.6, 5.5, 5.2, "🏆 Key Accomplishments", [
        "✅ Complete End-to-End System: Functional from physical sensors to actuation.",
        "✅ Real-World Standard Units: Temperature, Humidity, Lux, AQI, CO2, O2.",
        "✅ Botanical Intelligence: Advanced VPD modeling for transpiration health.",
        "✅ Configurable Thresholds: Live tuning sliders for any crop type.",
        "✅ Production-Ready Architecture: Flask REST API, SQLite, and Streamlit."
    ], color=GREEN)

    add_card(slide8, 6.8, 1.6, 5.7, 5.2, "🌱 Agricultural & Economic Value", [
        "Crop Protection: Prevents irreversible heat stress, frost, and fungal disease.",
        "Yield Maximization: Maintains optimal CO2 and VPD transpiration sweet-spots.",
        "Resource Conservation: Minimizes unnecessary fan/water pump power usage.",
        "Scalability: Easily expandable to multi-zone commercial greenhouse facilities.",
        "Thank You! Q&A Session."
    ], color=SLATE)

    output_path = os.path.abspath("Greenhouse_Monitoring_Presentation.pptx")
    prs.save(output_path)
    print(f"Presentation generated successfully: {output_path}")


if __name__ == "__main__":
    create_deck()
